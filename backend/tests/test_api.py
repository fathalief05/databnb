import csv
import io
import json
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.app.main import create_app


def make_client(tmp_path: Path) -> TestClient:
    return TestClient(create_app(tmp_path / "test.duckdb"))


def import_csv(client: TestClient) -> dict:
    response = client.post(
        "/api/datasets/import",
        files={"file": ("sales.csv", b"region,amount,region\nNorth,25,North\nSouth,60,South\n,90,\n", "text/csv")},
        data={"name": "Sales demo", "description": "Synthetic sample"},
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_preview_and_import_normalize_duplicate_headers_and_persist(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    preview = client.post(
        "/api/import/preview",
        files={"file": ("demo.csv", b"Name,Name\nAda,1\n", "text/csv")},
    )
    assert preview.status_code == 200
    assert [column["name"] for column in preview.json()["columns"]] == ["Name", "Name (2)"]

    dataset = import_csv(client)
    assert dataset["row_count"] == 3
    assert [column["name"] for column in dataset["columns"]] == [
        "region", "amount", "region (2)"
    ]
    reloaded = TestClient(create_app(tmp_path / "test.duckdb"))
    assert reloaded.get("/api/datasets").json()[0]["id"] == dataset["id"]
    assert reloaded.get(f"/api/datasets/{dataset['id']}/records").json()["total"] == 3


def test_xlsx_preview_and_import_select_a_worksheet(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "First"
    workbook.active.append(["ignored"])
    second = workbook.create_sheet("Inventory")
    second.append(["Item", "In stock"])
    second.append(["Desk", 7])
    file_content = BytesIO()
    workbook.save(file_content)
    client = make_client(tmp_path)

    preview = client.post(
        "/api/import/preview",
        files={"file": ("inventory.xlsx", file_content.getvalue(), "application/octet-stream")},
        data={"sheet_name": "Inventory"},
    )
    assert preview.status_code == 200
    assert preview.json()["sheets"] == ["First", "Inventory"]
    assert preview.json()["preview"] == [{"Item": "Desk", "In stock": 7}]

    imported = client.post(
        "/api/datasets/import",
        files={"file": ("inventory.xlsx", file_content.getvalue(), "application/octet-stream")},
        data={"name": "Inventory", "sheet_name": "Inventory"},
    )
    assert imported.status_code == 201, imported.text
    assert client.get(f"/api/datasets/{imported.json()['id']}/records").json()["records"][0]["In stock"] == 7

    bad_sheet = client.post(
        "/api/import/preview",
        files={"file": ("inventory.xlsx", file_content.getvalue(), "application/octet-stream")},
        data={"sheet_name": "Missing"},
    )
    assert bad_sheet.status_code == 422


def test_search_nested_filters_sort_and_pagination_use_full_dataset(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    dataset = import_csv(client)
    filters = {
        "logic": "or",
        "conditions": [
            {"column": "amount", "operator": "greater_than", "value": 80},
            {
                "logic": "and",
                "conditions": [
                    {"column": "region", "operator": "contains", "value": "nor"},
                    {"column": "amount", "operator": "less_than", "value": 30},
                ],
            },
        ],
    }
    response = client.get(
        f"/api/datasets/{dataset['id']}/records",
        params={
            "filters": json.dumps(filters),
            "sort_by": "amount",
            "sort_direction": "desc",
            "page_size": 1,
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["total"] == 2
    assert response.json()["records"][0]["amount"] == 90
    assert response.json()["pages"] == 2
    assert client.get(
        f"/api/datasets/{dataset['id']}/records", params={"search": "south"}
    ).json()["total"] == 1
    invalid_value = client.get(
        f"/api/datasets/{dataset['id']}/records",
        params={"filters": json.dumps({"column": "amount", "operator": "greater_than", "value": "many"})},
    )
    assert invalid_value.status_code == 422
    nulls = client.get(
        f"/api/datasets/{dataset['id']}/records",
        params={"filters": json.dumps({"column": "region", "operator": "is_null"})},
    )
    assert nulls.status_code == 200
    assert nulls.json()["total"] == 1


def test_date_columns_support_absolute_date_filters(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    imported = client.post(
        "/api/datasets/import",
        files={"file": ("dates.csv", b"day,created\n2025-01-05,2025-01-05T09:30:00\n2025-02-10,2025-02-10T14:00:00\n", "text/csv")},
        data={"name": "Date sample"},
    )
    assert imported.status_code == 201, imported.text
    dataset = imported.json()
    assert [column["type"] for column in dataset["columns"]] == ["DATE", "TIMESTAMP"]

    result = client.get(
        f"/api/datasets/{dataset['id']}/records",
        params={"filters": json.dumps({"column": "day", "operator": "after", "value": "2025-01-15"})},
    )
    assert result.status_code == 200, result.text
    assert result.json()["records"][0]["day"] == "2025-02-10"


def test_filters_reject_unknown_columns_and_sql_injection(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    dataset = import_csv(client)
    result = client.get(
        f"/api/datasets/{dataset['id']}/records",
        params={
            "filters": json.dumps(
                {
                    "column": "amount; DROP TABLE datasets",
                    "operator": "equals",
                    "value": "25 OR 1=1",
                }
            )
        },
    )
    assert result.status_code == 422
    assert client.get("/api/datasets").status_code == 200


def test_record_crud_configurations_and_export_scope(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    dataset = import_csv(client)
    record = client.get(f"/api/datasets/{dataset['id']}/records").json()["records"][0]
    record_id = record["_record_id"]

    update = client.patch(
        f"/api/datasets/{dataset['id']}/records/{record_id}",
        json={"values": {"amount": 35}},
    )
    assert update.status_code == 200
    assert update.json()["amount"] == 35
    created = client.post(
        f"/api/datasets/{dataset['id']}/records",
        json={"values": {"region": "East", "amount": 50, "region (2)": "East"}},
    )
    assert created.status_code == 201, created.text
    assert client.get(f"/api/datasets/{dataset['id']}/records").json()["total"] == 4
    assert client.delete(
        f"/api/datasets/{dataset['id']}/records/{created.json()['_record_id']}"
    ).status_code == 204

    config = client.post(
        f"/api/datasets/{dataset['id']}/configurations?kind=filter",
        json={
            "name": "Big orders",
            "description": "Large amounts",
            "definition": {"column": "amount", "operator": "greater_than", "value": 50},
            "schema_version": 1,
        },
    )
    assert config.status_code == 201, config.text
    configuration_id = config.json()["id"]
    assert client.get(
        f"/api/datasets/{dataset['id']}/configurations"
    ).json()[0]["id"] == configuration_id
    renamed = client.patch(
        f"/api/datasets/{dataset['id']}/configurations/{configuration_id}",
        json={
            "name": "High value",
            "description": "Renamed",
            "definition": {"column": "amount", "operator": "greater_than", "value": 50},
        },
    )
    assert renamed.status_code == 200
    assert renamed.json()["name"] == "High value"
    copied = client.post(
        f"/api/datasets/{dataset['id']}/configurations?kind=filter",
        json={
            "name": "High value copy",
            "definition": {"column": "amount", "operator": "greater_than", "value": 50},
        },
    )
    assert copied.status_code == 201
    restarted = make_client(tmp_path)
    assert len(restarted.get(
        f"/api/datasets/{dataset['id']}/configurations"
    ).json()) == 2

    exported = client.post(
        f"/api/datasets/{dataset['id']}/export",
        json={
            "format": "csv",
            "columns": ["region", "amount"],
            "filters": {"column": "amount", "operator": "greater_than", "value": 50},
        },
    )
    assert exported.status_code == 200
    rows = list(csv.reader(io.StringIO(exported.content.decode("utf-8-sig"))))
    assert rows == [["region", "amount"], ["South", "60"], ["", "90"]]
    excel_export = client.post(
        f"/api/datasets/{dataset['id']}/export",
        json={"format": "xlsx", "columns": ["amount"]},
    )
    assert excel_export.status_code == 200
    workbook = load_workbook(BytesIO(excel_export.content), read_only=True, data_only=True)
    assert list(workbook.active.values)[0] == ("amount",)
    workbook.close()
    assert client.delete(
        f"/api/datasets/{dataset['id']}/records/{record_id}"
    ).status_code == 204
    assert client.delete(
        f"/api/datasets/{dataset['id']}/configurations/{configuration_id}"
    ).status_code == 204
    assert client.delete(
        f"/api/datasets/{dataset['id']}/configurations/{copied.json()['id']}"
    ).status_code == 204


def test_dataset_deletion_removes_its_saved_configurations(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    dataset = import_csv(client)
    saved = client.post(
        f"/api/datasets/{dataset['id']}/configurations?kind=table",
        json={"name": "Compact", "definition": {"visibleColumns": ["region"]}},
    )
    assert saved.status_code == 201
    assert client.delete(f"/api/datasets/{dataset['id']}").status_code == 204
    assert client.get(f"/api/datasets/{dataset['id']}").status_code == 404
    assert client.get(f"/api/datasets/{dataset['id']}/configurations").status_code == 404


def test_saved_configuration_definitions_are_validated(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    dataset = import_csv(client)
    invalid_filter = client.post(
        f"/api/datasets/{dataset['id']}/configurations?kind=filter",
        json={
            "name": "Unsafe",
            "definition": {
                "column": "amount; DROP TABLE datasets",
                "operator": "equals",
                "value": 1,
            },
        },
    )
    invalid_table = client.post(
        f"/api/datasets/{dataset['id']}/configurations?kind=table",
        json={"name": "Missing field", "definition": {"visibleColumns": ["missing"]}},
    )
    assert invalid_filter.status_code == 422
    assert invalid_table.status_code == 422


def test_csv_preview_falls_back_to_windows_1252(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    response = client.post(
        "/api/import/preview",
        files={"file": ("names.csv", b"name\nAndr\xe9\n", "text/csv")},
    )
    assert response.status_code == 200
    assert response.json()["preview"] == [{"name": "André"}]


def test_uploaded_column_names_are_quoted_as_display_labels(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    response = client.post(
        "/api/datasets/import",
        files={"file": ("quoted.csv", b'"column ""quoted""",value\nsafe,1\n', "text/csv")},
        data={"name": "Quoted headers"},
    )
    assert response.status_code == 201, response.text
    dataset = response.json()
    record = client.get(f"/api/datasets/{dataset['id']}/records").json()["records"][0]
    assert record['column "quoted"'] == "safe"
    assert client.get("/api/datasets").status_code == 200


def test_reject_oversized_and_malformed_files(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    malformed = client.post(
        "/api/import/preview",
        files={"file": ("broken.xlsx", b"not excel", "application/octet-stream")},
    )
    assert malformed.status_code == 422
    oversized = client.post(
        "/api/import/preview",
        files={"file": ("big.csv", b"x" * (25 * 1024 * 1024 + 1), "text/csv")},
    )
    assert oversized.status_code == 413
