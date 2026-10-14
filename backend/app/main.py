from __future__ import annotations

import csv
import io
import json
import os
import re
import threading
import uuid
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import duckdb
from fastapi import FastAPI, File, Form, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field, field_validator

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 200
MAX_CELLS = 2_000_000
MAX_UNCOMPRESSED_XLSX_BYTES = 200 * 1024 * 1024
MAX_PAGE_SIZE = 200


def quote_identifier(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


def normalize_headers(headers: list[Any], width: int) -> list[str]:
    names: list[str] = []
    used: set[str] = set()
    for index in range(width):
        value = headers[index] if index < len(headers) else ""
        name = str(value).strip() if value is not None else ""
        name = name or f"Column {index + 1}"
        candidate = name
        suffix = 2
        while candidate.casefold() in used:
            candidate = f"{name} ({suffix})"
            suffix += 1
        used.add(candidate.casefold())
        names.append(candidate)
    return names


def upload_filename(filename: str) -> str:
    return filename.replace("\\", "/").rsplit("/", 1)[-1] or "upload"


def read_spreadsheet(
    content: bytes, filename: str, sheet_name: str | None = None
) -> tuple[list[str], list[list[Any]], list[str]]:
    extension = Path(filename).suffix.casefold()
    if extension == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("cp1252")
            except UnicodeDecodeError as exc:
                raise ValueError("Could not decode this CSV file") from exc
        rows = []
        for row in csv.reader(io.StringIO(text, newline="")):
            if len(rows) >= MAX_ROWS + 1:
                raise ValueError(f"Files may contain at most {MAX_ROWS:,} records")
            if len(row) > MAX_COLUMNS:
                raise ValueError(f"Files may contain at most {MAX_COLUMNS} columns")
            if len(rows) * max(1, len(row)) > MAX_CELLS:
                raise ValueError("This file contains too many cells to import safely")
            rows.append(row)
        available_sheets: list[str] = []
    elif extension == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_XLSX_BYTES:
                    raise ValueError("This Excel file expands beyond the supported size")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("This Excel file is malformed or cannot be opened") from exc
        available_sheets = workbook.sheetnames
        selected = sheet_name or (available_sheets[0] if available_sheets else None)
        if not selected or selected not in available_sheets:
            raise ValueError("Select a worksheet that exists in this workbook")
        worksheet = workbook[selected]
        if worksheet.max_row and worksheet.max_row > MAX_ROWS + 1:
            workbook.close()
            raise ValueError(f"Files may contain at most {MAX_ROWS:,} records")
        if worksheet.max_column and worksheet.max_column > MAX_COLUMNS:
            workbook.close()
            raise ValueError(f"Files may contain at most {MAX_COLUMNS} columns")
        rows = []
        for row in worksheet.iter_rows(values_only=True):
            if len(rows) >= MAX_ROWS + 1:
                workbook.close()
                raise ValueError(f"Files may contain at most {MAX_ROWS:,} records")
            if len(rows) * max(1, len(row)) > MAX_CELLS:
                workbook.close()
                raise ValueError("This file contains too many cells to import safely")
            rows.append(list(row))
        workbook.close()
    else:
        raise ValueError("Only CSV and XLSX files are supported")
    while rows and not any(value not in (None, "") for value in rows[-1]):
        rows.pop()
    if not rows:
        raise ValueError("The file does not contain a header row")
    width = max((len(row) for row in rows), default=0)
    headers = normalize_headers(rows[0], width)
    records = [
        (list(row) + [None] * width)[:width]
        for row in rows[1:]
    ]
    return headers, records, available_sheets


def infer_column_type(values: list[Any]) -> str:
    nonempty = [value for value in values if value is not None and str(value).strip() != ""]
    if not nonempty:
        return "VARCHAR"
    if all(isinstance(value, datetime) for value in nonempty):
        return "TIMESTAMP"
    if all(isinstance(value, date) and not isinstance(value, datetime) for value in nonempty):
        return "DATE"
    strings = [str(value).strip() for value in nonempty]
    if all(value.casefold() in {"true", "false"} for value in strings):
        return "BOOLEAN"
    if all(re.fullmatch(r"[+-]?\d+", value) for value in strings):
        try:
            if all(-(2**63) <= int(value) < 2**63 for value in strings):
                return "BIGINT"
        except ValueError:
            pass
    if all(_is_number(value) for value in strings):
        return "DOUBLE"
    try:
        if all(date.fromisoformat(value) for value in strings):
            return "DATE"
    except ValueError:
        pass
    try:
        if all(datetime.fromisoformat(value.replace("Z", "+00:00")) for value in strings):
            return "TIMESTAMP"
    except ValueError:
        pass
    return "VARCHAR"


def _is_number(value: str) -> bool:
    try:
        float(value)
        return True
    except ValueError:
        return False


def _coerce(value: Any, column_type: str) -> Any:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, (date, datetime, time)):
        value = value.isoformat()
    if column_type == "BOOLEAN":
        text = str(value).strip().casefold()
        if text not in {"true", "false"}:
            raise ValueError(f"Expected true or false, received {value!r}")
        return text == "true"
    if column_type == "BIGINT":
        return int(value)
    if column_type == "DOUBLE":
        return float(value)
    if column_type == "DATE":
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        return date.fromisoformat(str(value))
    if column_type == "TIMESTAMP":
        if isinstance(value, datetime):
            return value
        return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    return str(value)


def _table_name(dataset_id: str) -> str:
    try:
        parsed = uuid.UUID(dataset_id)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail="Dataset not found") from exc
    return f"dataset_{parsed.hex}"


def _column_names(dataset: dict[str, Any]) -> list[str]:
    return [column["name"] for column in dataset["columns"]]


def _physical_column(dataset: dict[str, Any], name: str) -> str:
    for index, column in enumerate(dataset["columns"]):
        if column["name"] == name:
            return f"c{index}"
    raise HTTPException(status_code=422, detail=f"Unknown column: {name}")


def _compile_filter(
    dataset: dict[str, Any], node: Any, depth: int = 0
) -> tuple[str, list[Any]]:
    if node is None:
        return "", []
    if depth > 5 or not isinstance(node, dict):
        raise HTTPException(status_code=422, detail="Invalid filter definition")
    if "conditions" in node:
        conditions = node["conditions"]
        logic = str(node.get("logic", "and")).casefold()
        if (
            logic not in {"and", "or"}
            or not isinstance(conditions, list)
            or len(conditions) > 50
        ):
            raise HTTPException(status_code=422, detail="Invalid filter group")
        compiled = [
            _compile_filter(dataset, condition, depth + 1)
            for condition in conditions
        ]
        compiled = [(sql, params) for sql, params in compiled if sql]
        if not compiled:
            return "", []
        return (
            "(" + f" {logic.upper()} ".join(sql for sql, _ in compiled) + ")",
            [value for _, params in compiled for value in params],
        )

    name = node.get("column")
    if not isinstance(name, str):
        raise HTTPException(status_code=422, detail="A filter column is required")
    column = quote_identifier(_physical_column(dataset, name))
    column_type = dataset["columns"][
        next(index for index, item in enumerate(dataset["columns"]) if item["name"] == name)
    ]["type"]
    operator = str(node.get("operator", "")).casefold()
    value = node.get("value")

    def typed_filter_value(raw: Any) -> Any:
        if raw is not None and not isinstance(raw, (str, int, float, bool)):
            raise HTTPException(status_code=422, detail="Filter values must be simple values")
        try:
            return _coerce(raw, column_type)
        except (ValueError, TypeError, OverflowError) as exc:
            raise HTTPException(
                status_code=422,
                detail=f"Invalid {column_type.lower()} filter value",
            ) from exc

    if operator == "is_null":
        return f"{column} IS NULL", []
    if operator == "is_not_null":
        return f"{column} IS NOT NULL", []
    if operator in {"is_empty", "is_not_empty"}:
        comparison = "=" if operator == "is_empty" else "<>"
        return f"({column} IS NOT NULL AND CAST({column} AS VARCHAR) {comparison} '')", []
    if operator in {"true", "false"}:
        return f"{column} = ?", [operator == "true"]
    if operator in {"contains", "does_not_contain"}:
        if value is not None and not isinstance(value, (str, int, float, bool)):
            raise HTTPException(status_code=422, detail="Filter values must be simple values")
        expression = f"contains(lower(CAST({column} AS VARCHAR)), lower(?))"
        return (
            f"NOT {expression}" if operator == "does_not_contain" else expression,
            [str(value if value is not None else "")],
        )
    if operator in {"starts_with", "ends_with"}:
        function = "starts_with" if operator == "starts_with" else "ends_with"
        return f"{function}(lower(CAST({column} AS VARCHAR)), lower(?))", [str(value or "")]
    if operator in {"on", "before", "after"}:
        column_type = dataset["columns"][
            next(i for i, col in enumerate(dataset["columns"]) if col["name"] == name)
        ]["type"]
        if column_type not in {"DATE", "TIMESTAMP"}:
            raise HTTPException(status_code=422, detail="Date comparisons require a date column")
        comparison = {"on": "=", "before": "<", "after": ">"}[operator]
        return f"{column} {comparison} ?", [typed_filter_value(value)]
    if operator in {"in", "not_in"}:
        if (
            not isinstance(value, list)
            or len(value) > 100
            or any(not isinstance(item, (str, int, float, bool)) for item in value)
        ):
            raise HTTPException(status_code=422, detail="Select up to 100 values")
        if not value:
            return ("FALSE" if operator == "in" else "TRUE"), []
        placeholders = ", ".join("?" for _ in value)
        keyword = "NOT IN" if operator == "not_in" else "IN"
        return f"{column} {keyword} ({placeholders})", [typed_filter_value(item) for item in value]
    if operator == "between":
        if (
            not isinstance(value, list)
            or len(value) != 2
            or any(not isinstance(item, (str, int, float, bool)) for item in value)
        ):
            raise HTTPException(status_code=422, detail="Between requires two values")
        return f"{column} BETWEEN ? AND ?", [typed_filter_value(item) for item in value]
    operators = {
        "equals": "=",
        "on": "=",
        "not_equals": "<>",
        "before": "<",
        "after": ">",
        "greater_than": ">",
        "greater_than_or_equal": ">=",
        "less_than": "<",
        "less_than_or_equal": "<=",
    }
    sql_operator = operators.get(operator)
    if not sql_operator or isinstance(value, (dict, list)):
        raise HTTPException(status_code=422, detail="Unsupported filter operator or value")
    if operator in {"greater_than", "greater_than_or_equal", "less_than", "less_than_or_equal"}:
        if column_type not in {"BIGINT", "DOUBLE"}:
            raise HTTPException(status_code=422, detail="Numeric comparisons require a numeric column")
    return f"{column} {sql_operator} ?", [typed_filter_value(value)]


def _validate_configuration_definition(
    dataset: dict[str, Any], kind: str, definition: dict[str, Any]
) -> None:
    columns = set(_column_names(dataset))
    if kind == "filter":
        _compile_filter(dataset, definition)
        return
    if kind == "table":
        visible = definition.get("visibleColumns", [])
        sort_by = definition.get("sortBy")
        direction = definition.get("sortDirection", "asc")
        if (
            not isinstance(visible, list)
            or any(not isinstance(name, str) or name not in columns for name in visible)
            or len(visible) > len(columns)
            or len(visible) != len(set(visible))
        ):
            raise HTTPException(status_code=422, detail="Table view contains invalid columns")
        if sort_by is not None and (not isinstance(sort_by, str) or sort_by not in columns):
            raise HTTPException(status_code=422, detail="Table view contains an invalid sort column")
        if not isinstance(direction, str) or direction not in {"asc", "desc"}:
            raise HTTPException(status_code=422, detail="Table view sort direction must be asc or desc")
        return
    if kind == "detail":
        fields = definition.get("fields", [])
        labels = definition.get("labels", {})
        section = definition.get("section", "Record details")
        column_count = definition.get("columns", 2)
        if (
            not isinstance(fields, list)
            or any(not isinstance(name, str) or name not in columns for name in fields)
            or len(fields) > len(columns)
            or len(fields) != len(set(fields))
        ):
            raise HTTPException(status_code=422, detail="Detail view contains invalid fields")
        if (
            not isinstance(labels, dict)
            or any(
                not isinstance(name, str)
                or name not in columns
                or not isinstance(label, str)
                or len(label) > 120
                for name, label in labels.items()
            )
        ):
            raise HTTPException(status_code=422, detail="Detail view contains invalid display labels")
        if (
            not isinstance(section, str)
            or len(section) > 120
            or type(column_count) is not int
            or column_count not in {1, 2}
        ):
            raise HTTPException(status_code=422, detail="Detail view presentation is invalid")
        return
    raise HTTPException(status_code=422, detail="Unknown configuration type")


class RecordInput(BaseModel):
    values: dict[str, Any] = Field(default_factory=dict)


class ConfigurationInput(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    description: str | None = Field(default=None, max_length=500)
    definition: dict[str, Any]
    schema_version: int = Field(default=1, ge=1)

    @field_validator("name")
    @classmethod
    def name_must_not_be_blank(cls, value: str) -> str:
        if not value.strip():
            raise ValueError("Name cannot be blank")
        return value.strip()


class Database:
    def __init__(self, path: str | Path):
        self.path = str(path)
        Path(self.path).parent.mkdir(parents=True, exist_ok=True)
        self.connection = duckdb.connect(self.path)
        self.lock = threading.RLock()
        with self.lock:
            self.connection.execute(
                """CREATE TABLE IF NOT EXISTS datasets (
                    id VARCHAR PRIMARY KEY,
                    name VARCHAR NOT NULL,
                    description VARCHAR NOT NULL,
                    source_filename VARCHAR NOT NULL,
                    source_type VARCHAR NOT NULL,
                    row_count BIGINT NOT NULL,
                    columns_json VARCHAR NOT NULL,
                    created_at TIMESTAMP NOT NULL DEFAULT current_timestamp
                )"""
            )
            self.connection.execute(
                """CREATE TABLE IF NOT EXISTS saved_configurations (
                    id VARCHAR PRIMARY KEY,
                    dataset_id VARCHAR NOT NULL,
                    kind VARCHAR NOT NULL,
                    name VARCHAR NOT NULL,
                    description VARCHAR NOT NULL,
                    definition_json VARCHAR NOT NULL,
                    schema_version INTEGER NOT NULL,
                    created_at TIMESTAMP NOT NULL DEFAULT current_timestamp,
                    updated_at TIMESTAMP NOT NULL DEFAULT current_timestamp,
                    UNIQUE(dataset_id, kind, name)
                )"""
            )

    def dataset(self, dataset_id: str) -> dict[str, Any]:
        with self.lock:
            row = self.connection.execute(
                "SELECT id, name, description, source_filename, source_type, row_count, columns_json, created_at FROM datasets WHERE id = ?",
                [dataset_id],
            ).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Dataset not found")
        return {
            "id": row[0],
            "name": row[1],
            "description": row[2],
            "source_filename": row[3],
            "source_type": row[4],
            "row_count": row[5],
            "columns": json.loads(row[6]),
            "created_at": row[7].isoformat(),
        }


def create_app(database_path: str | Path | None = None) -> FastAPI:
    configured_path = database_path or os.environ.get(
        "DATABNB_DATABASE_PATH", "data/databnb.duckdb"
    )
    database = Database(configured_path)
    app = FastAPI(title="DataBNB", version="0.1.0")
    app.state.database = database
    app.add_middleware(
        CORSMiddleware,
        allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
        allow_credentials=False,
        allow_methods=["*"],
        allow_headers=["*"],
    )

    @app.get("/api/health")
    def health() -> dict[str, str]:
        return {"status": "ok"}

    @app.get("/api/datasets")
    def list_datasets() -> list[dict[str, Any]]:
        with database.lock:
            rows = database.connection.execute(
                "SELECT id, name, description, source_filename, source_type, row_count, columns_json, created_at FROM datasets ORDER BY created_at DESC"
            ).fetchall()
        return [
            {
                "id": row[0],
                "name": row[1],
                "description": row[2],
                "source_filename": row[3],
                "source_type": row[4],
                "row_count": row[5],
                "columns": json.loads(row[6]),
                "created_at": row[7].isoformat(),
            }
            for row in rows
        ]

    @app.post("/api/import/preview")
    async def preview_import(
        file: UploadFile = File(...), sheet_name: str | None = Form(default=None)
    ) -> dict[str, Any]:
        content = await file.read(MAX_UPLOAD_BYTES + 1)
        if len(content) > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=413, detail="Files must be 25 MB or smaller")
        try:
            headers, rows, sheets = read_spreadsheet(content, file.filename or "", sheet_name)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        columns = [
            {"name": name, "type": infer_column_type([row[index] for row in rows])}
            for index, name in enumerate(headers)
        ]
        return {
            "filename": upload_filename(file.filename or ""),
            "sheets": sheets,
            "columns": columns,
            "row_count": len(rows),
            "preview": [dict(zip(headers, row)) for row in rows[:20]],
        }

    @app.post("/api/datasets/import", status_code=201)
    async def import_dataset(
        file: UploadFile = File(...),
        name: str = Form(min_length=1, max_length=120),
        description: str = Form(default="", max_length=500),
        sheet_name: str | None = Form(default=None),
    ) -> dict[str, Any]:
        content = await file.read(MAX_UPLOAD_BYTES + 1)
        if len(content) > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=413, detail="Files must be 25 MB or smaller")
        if not name.strip():
            raise HTTPException(status_code=422, detail="Dataset name cannot be blank")
        try:
            headers, rows, _ = read_spreadsheet(content, file.filename or "", sheet_name)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        columns = [
            {"name": header, "type": infer_column_type([row[index] for row in rows])}
            for index, header in enumerate(headers)
        ]
        dataset_id = str(uuid.uuid4())
        table = _table_name(dataset_id)
        table_sql = quote_identifier(table)
        definitions = ['"_record_id" UUID PRIMARY KEY'] + [
            f'{quote_identifier(f"c{index}")} {column["type"]}'
            for index, column in enumerate(columns)
        ]
        with database.lock:
            connection = database.connection
            try:
                connection.execute("BEGIN TRANSACTION")
                connection.execute(f"CREATE TABLE {table_sql} ({', '.join(definitions)})")
                if rows:
                    placeholders = ", ".join("?" for _ in definitions)
                    insert = f"INSERT INTO {table_sql} VALUES ({placeholders})"
                    values = [
                        [str(uuid.uuid4())]
                        + [
                            _coerce(row[index], columns[index]["type"])
                            for index in range(len(columns))
                        ]
                        for row in rows
                    ]
                    connection.executemany(insert, values)
                connection.execute(
                    "INSERT INTO datasets (id, name, description, source_filename, source_type, row_count, columns_json) VALUES (?, ?, ?, ?, ?, ?, ?)",
                    [
                        dataset_id,
                        name.strip(),
                        description.strip(),
                        upload_filename(file.filename or "upload"),
                        Path(file.filename or "").suffix.casefold().lstrip("."),
                        len(rows),
                        json.dumps(columns),
                    ],
                )
                connection.execute("COMMIT")
            except Exception as exc:
                connection.execute("ROLLBACK")
                raise HTTPException(status_code=422, detail=f"Could not import file: {exc}") from exc
        return database.dataset(dataset_id)

    @app.get("/api/datasets/{dataset_id}")
    def get_dataset(dataset_id: str) -> dict[str, Any]:
        return database.dataset(dataset_id)

    @app.get("/api/datasets/{dataset_id}/records")
    def list_records(
        dataset_id: str,
        search: str = "",
        page: int = Query(default=1, ge=1),
        page_size: int = Query(default=50, ge=1, le=MAX_PAGE_SIZE),
        sort_by: str | None = None,
        sort_direction: str = "asc",
        filters: str | None = None,
    ) -> dict[str, Any]:
        dataset = database.dataset(dataset_id)
        table = quote_identifier(_table_name(dataset_id))
        where: list[str] = []
        parameters: list[Any] = []
        if search.strip() and dataset["columns"]:
            searchable = [
                f"contains(lower(CAST({quote_identifier(f'c{index}')} AS VARCHAR)), lower(?))"
                for index in range(len(dataset["columns"]))
            ]
            where.append("(" + " OR ".join(searchable) + ")")
            parameters.extend([search.strip()] * len(searchable))
        if filters:
            try:
                filter_definition = json.loads(filters)
            except json.JSONDecodeError as exc:
                raise HTTPException(status_code=422, detail="Filters must be valid JSON") from exc
            expression, values = _compile_filter(dataset, filter_definition)
            if expression:
                where.append(expression)
                parameters.extend(values)
        where_sql = " WHERE " + " AND ".join(where) if where else ""
        if sort_by is not None:
            physical = quote_identifier(_physical_column(dataset, sort_by))
            direction = sort_direction.casefold()
            if direction not in {"asc", "desc"}:
                raise HTTPException(status_code=422, detail="Sort direction must be asc or desc")
            order_sql = f" ORDER BY {physical} {direction.upper()} NULLS LAST"
        else:
            order_sql = ""
        projection = [
            '"_record_id"'
        ] + [
            f'{quote_identifier(f"c{index}")} AS {quote_identifier(column["name"])}'
            for index, column in enumerate(dataset["columns"])
        ]
        query_from = f" FROM {table}{where_sql}"
        with database.lock:
            total = database.connection.execute(
                "SELECT count(*)" + query_from, parameters
            ).fetchone()[0]
            result = database.connection.execute(
                "SELECT " + ", ".join(projection) + query_from + order_sql + " LIMIT ? OFFSET ?",
                [*parameters, page_size, (page - 1) * page_size],
            )
            names = [column[0] for column in result.description]
            records = [dict(zip(names, row)) for row in result.fetchall()]
        return {
            "records": records,
            "total": total,
            "page": page,
            "page_size": page_size,
            "pages": (total + page_size - 1) // page_size,
        }

    @app.get("/api/datasets/{dataset_id}/records/{record_id}")
    def get_record(dataset_id: str, record_id: str) -> dict[str, Any]:
        dataset = database.dataset(dataset_id)
        table = quote_identifier(_table_name(dataset_id))
        projection = [
            '"_record_id"'
        ] + [
            f'{quote_identifier(f"c{index}")} AS {quote_identifier(column["name"])}'
            for index, column in enumerate(dataset["columns"])
        ]
        with database.lock:
            row = database.connection.execute(
                "SELECT " + ", ".join(projection) + f" FROM {table} WHERE _record_id = ?",
                [record_id],
            ).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Record not found")
        return dict(zip(["_record_id", *_column_names(dataset)], row))

    @app.post("/api/datasets/{dataset_id}/records", status_code=201)
    def create_record(dataset_id: str, payload: RecordInput) -> dict[str, Any]:
        dataset = database.dataset(dataset_id)
        unknown = set(payload.values) - set(_column_names(dataset))
        if unknown:
            raise HTTPException(status_code=422, detail=f"Unknown columns: {', '.join(sorted(unknown))}")
        record_id = str(uuid.uuid4())
        physical = [
            f"c{index}" for index, column in enumerate(dataset["columns"])
        ]
        columns = ['"_record_id"', *[quote_identifier(name) for name in physical]]
        values = [record_id]
        try:
            values.extend(
                _coerce(
                    payload.values.get(column["name"]),
                    column["type"],
                )
                for column in dataset["columns"]
            )
        except (ValueError, TypeError) as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        table = quote_identifier(_table_name(dataset_id))
        with database.lock:
            database.connection.execute(
                f"INSERT INTO {table} ({', '.join(columns)}) VALUES ({', '.join('?' for _ in columns)})",
                values,
            )
            database.connection.execute(
                "UPDATE datasets SET row_count = row_count + 1 WHERE id = ?", [dataset_id]
            )
        return get_record(dataset_id, record_id)

    @app.patch("/api/datasets/{dataset_id}/records/{record_id}")
    def update_record(
        dataset_id: str, record_id: str, payload: RecordInput
    ) -> dict[str, Any]:
        dataset = database.dataset(dataset_id)
        unknown = set(payload.values) - set(_column_names(dataset))
        if unknown:
            raise HTTPException(status_code=422, detail=f"Unknown columns: {', '.join(sorted(unknown))}")
        if not payload.values:
            return get_record(dataset_id, record_id)
        assignments: list[str] = []
        values: list[Any] = []
        try:
            for index, column in enumerate(dataset["columns"]):
                if column["name"] in payload.values:
                    assignments.append(f"{quote_identifier(f'c{index}')} = ?")
                    values.append(_coerce(payload.values[column["name"]], column["type"]))
        except (ValueError, TypeError) as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        table = quote_identifier(_table_name(dataset_id))
        with database.lock:
            result = database.connection.execute(
                f"UPDATE {table} SET {', '.join(assignments)} WHERE _record_id = ? RETURNING _record_id",
                [*values, record_id],
            ).fetchone()
        if result is None:
            raise HTTPException(status_code=404, detail="Record not found")
        return get_record(dataset_id, record_id)

    @app.delete("/api/datasets/{dataset_id}/records/{record_id}", status_code=204)
    def delete_record(dataset_id: str, record_id: str) -> None:
        database.dataset(dataset_id)
        table = quote_identifier(_table_name(dataset_id))
        with database.lock:
            result = database.connection.execute(
                f"DELETE FROM {table} WHERE _record_id = ? RETURNING _record_id", [record_id]
            ).fetchone()
            if result is None:
                raise HTTPException(status_code=404, detail="Record not found")
            database.connection.execute(
                "UPDATE datasets SET row_count = row_count - 1 WHERE id = ?", [dataset_id]
            )

    @app.post("/api/datasets/{dataset_id}/configurations", status_code=201)
    def save_configuration(
        dataset_id: str,
        kind: str = Query(pattern="^(filter|table|detail)$"),
        payload: ConfigurationInput = ...,
    ) -> dict[str, Any]:
        dataset = database.dataset(dataset_id)
        _validate_configuration_definition(dataset, kind, payload.definition)
        configuration_id = str(uuid.uuid4())
        try:
            with database.lock:
                database.connection.execute(
                    "INSERT INTO saved_configurations (id, dataset_id, kind, name, description, definition_json, schema_version) VALUES (?, ?, ?, ?, ?, ?, ?)",
                    [
                        configuration_id,
                        dataset_id,
                        kind,
                        payload.name.strip(),
                        (payload.description or "").strip(),
                        json.dumps(payload.definition),
                        payload.schema_version,
                    ],
                )
        except duckdb.ConstraintException as exc:
            raise HTTPException(status_code=409, detail="A configuration with this name already exists") from exc
        return get_configuration(dataset_id, configuration_id)

    def configuration_dict(row: tuple[Any, ...]) -> dict[str, Any]:
        return {
            "id": row[0],
            "dataset_id": row[1],
            "kind": row[2],
            "name": row[3],
            "description": row[4],
            "definition": json.loads(row[5]),
            "schema_version": row[6],
            "created_at": row[7].isoformat(),
            "updated_at": row[8].isoformat(),
        }

    def get_configuration(dataset_id: str, configuration_id: str) -> dict[str, Any]:
        with database.lock:
            row = database.connection.execute(
                "SELECT id, dataset_id, kind, name, description, definition_json, schema_version, created_at, updated_at FROM saved_configurations WHERE dataset_id = ? AND id = ?",
                [dataset_id, configuration_id],
            ).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Configuration not found")
        return configuration_dict(row)

    @app.get("/api/datasets/{dataset_id}/configurations")
    def list_configurations(
        dataset_id: str, kind: str | None = Query(default=None, pattern="^(filter|table|detail)$")
    ) -> list[dict[str, Any]]:
        database.dataset(dataset_id)
        sql = "SELECT id, dataset_id, kind, name, description, definition_json, schema_version, created_at, updated_at FROM saved_configurations WHERE dataset_id = ?"
        parameters: list[Any] = [dataset_id]
        if kind:
            sql += " AND kind = ?"
            parameters.append(kind)
        with database.lock:
            rows = database.connection.execute(sql + " ORDER BY name", parameters).fetchall()
        return [configuration_dict(row) for row in rows]

    @app.patch("/api/datasets/{dataset_id}/configurations/{configuration_id}")
    def update_configuration(
        dataset_id: str, configuration_id: str, payload: ConfigurationInput
    ) -> dict[str, Any]:
        existing = get_configuration(dataset_id, configuration_id)
        _validate_configuration_definition(
            database.dataset(dataset_id), existing["kind"], payload.definition
        )
        try:
            with database.lock:
                database.connection.execute(
                    "UPDATE saved_configurations SET name = ?, description = ?, definition_json = ?, schema_version = ?, updated_at = current_timestamp WHERE dataset_id = ? AND id = ?",
                    [
                        payload.name.strip(),
                        (payload.description or "").strip(),
                        json.dumps(payload.definition),
                        payload.schema_version,
                        dataset_id,
                        configuration_id,
                    ],
                )
        except duckdb.ConstraintException as exc:
            raise HTTPException(status_code=409, detail="A configuration with this name already exists") from exc
        return get_configuration(dataset_id, configuration_id)

    @app.delete("/api/datasets/{dataset_id}/configurations/{configuration_id}", status_code=204)
    def delete_configuration(dataset_id: str, configuration_id: str) -> None:
        with database.lock:
            result = database.connection.execute(
                "DELETE FROM saved_configurations WHERE dataset_id = ? AND id = ? RETURNING id",
                [dataset_id, configuration_id],
            ).fetchone()
        if result is None:
            raise HTTPException(status_code=404, detail="Configuration not found")

    @app.post("/api/datasets/{dataset_id}/export")
    def export_dataset(dataset_id: str, payload: dict[str, Any]) -> StreamingResponse:
        dataset = database.dataset(dataset_id)
        output_format = str(payload.get("format", "csv")).casefold()
        if output_format not in {"csv", "xlsx"}:
            raise HTTPException(status_code=422, detail="Export format must be CSV or XLSX")
        selected = payload.get("columns", _column_names(dataset))
        if not isinstance(selected, list) or not all(
            isinstance(name, str) and name in _column_names(dataset) for name in selected
        ):
            raise HTTPException(status_code=422, detail="Invalid export columns")
        expression, parameters = _compile_filter(dataset, payload.get("filters"))
        where_sql = f" WHERE {expression}" if expression else ""
        if payload.get("search"):
            searchable = [
                f"contains(lower(CAST({quote_identifier(f'c{index}')} AS VARCHAR)), lower(?))"
                for index in range(len(dataset["columns"]))
            ]
            if searchable:
                where_sql += (" AND " if where_sql else " WHERE ") + "(" + " OR ".join(searchable) + ")"
                parameters.extend([str(payload["search"]).strip()] * len(searchable))
        column_indexes = {
            column["name"]: index for index, column in enumerate(dataset["columns"])
        }
        select_parts = [
            f'{quote_identifier(f"c{column_indexes[name]}")} AS {quote_identifier(name)}'
            for name in selected
        ]
        query = (
            "SELECT " + (", ".join(select_parts) if select_parts else "1 AS empty")
            + f" FROM {quote_identifier(_table_name(dataset_id))}{where_sql}"
        )
        with database.lock:
            rows = database.connection.execute(query, parameters).fetchall()
        if output_format == "csv":
            text = io.StringIO(newline="")
            writer = csv.writer(text)
            writer.writerow(selected)
            writer.writerows(rows if selected else [])
            content = text.getvalue().encode("utf-8-sig")
            media_type = "text/csv; charset=utf-8"
        else:
            workbook = Workbook(write_only=True)
            worksheet = workbook.create_sheet("Data")
            worksheet.append(selected)
            for row in rows:
                worksheet.append(list(row))
            stream = io.BytesIO()
            workbook.save(stream)
            content = stream.getvalue()
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        return StreamingResponse(
            io.BytesIO(content),
            media_type=media_type,
            headers={
                "Content-Disposition": (
                    "attachment; filename=\""
                    + (re.sub(r"[^A-Za-z0-9._-]+", "_", dataset["name"]).strip("._") or "dataset")
                    + f".{output_format}\""
                )
            },
        )

    @app.delete("/api/datasets/{dataset_id}", status_code=204)
    def delete_dataset(dataset_id: str) -> None:
        dataset = database.dataset(dataset_id)
        table = quote_identifier(_table_name(dataset_id))
        with database.lock:
            database.connection.execute("BEGIN TRANSACTION")
            try:
                database.connection.execute(f"DROP TABLE {table}")
                database.connection.execute(
                    "DELETE FROM saved_configurations WHERE dataset_id = ?", [dataset_id]
                )
                database.connection.execute("DELETE FROM datasets WHERE id = ?", [dataset_id])
                database.connection.execute("COMMIT")
            except Exception:
                database.connection.execute("ROLLBACK")
                raise

    return app


app = create_app()
