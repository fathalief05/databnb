import csv
import io
import json
import math
import uuid
from datetime import date, datetime
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook

from app.connectors.spreadsheets import infer_schema
from app.core.database import DB_LOCK, connect, decode_json, quote_identifier

DUCK_TYPES = {
    "text": "VARCHAR",
    "integer": "BIGINT",
    "number": "DOUBLE",
    "boolean": "BOOLEAN",
    "date": "DATE",
    "datetime": "TIMESTAMP",
}
FILTER_OPERATORS = {
    "text": {"equals", "not_equals", "contains", "not_contains", "starts_with", "ends_with", "is_empty", "is_not_empty", "is_null", "is_not_null"},
    "integer": {"equals", "not_equals", "greater_than", "greater_or_equal", "less_than", "less_or_equal", "between", "is_null", "is_not_null"},
    "number": {"equals", "not_equals", "greater_than", "greater_or_equal", "less_than", "less_or_equal", "between", "is_null", "is_not_null"},
    "date": {"on", "before", "after", "between", "is_null", "is_not_null"},
    "datetime": {"on", "before", "after", "between", "is_null", "is_not_null"},
    "boolean": {"equals", "is_null", "is_not_null"},
}


def _dataset(connection, dataset_id: str):
    row = connection.execute(
        "SELECT id, name, description, source_type, source_filename, table_name, schema_json, row_count, created_at, updated_at "
        "FROM _datasets WHERE id = ?",
        [dataset_id],
    ).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="Dataset not found.")
    keys = ("id", "name", "description", "source_type", "source_filename", "table_name", "schema_json", "row_count", "created_at", "updated_at")
    dataset = dict(zip(keys, row))
    dataset["schema"] = decode_json(dataset.pop("schema_json"))
    return dataset


def list_datasets():
    with DB_LOCK:
        connection = connect()
        try:
            rows = connection.execute(
                "SELECT id, name, description, source_type, source_filename, table_name, schema_json, row_count, created_at, updated_at "
                "FROM _datasets ORDER BY created_at DESC"
            ).fetchall()
            keys = ("id", "name", "description", "source_type", "source_filename", "table_name", "schema_json", "row_count", "created_at", "updated_at")
            result = []
            for row in rows:
                item = dict(zip(keys, row))
                item["schema"] = decode_json(item.pop("schema_json"))
                item.pop("table_name")
                result.append(item)
            return result
        finally:
            connection.close()


def import_dataset(name: str, description: str, filename: str, headers: list[str], rows: list[list[Any]], source_type: str):
    if not headers:
        raise HTTPException(status_code=400, detail="The selected sheet does not contain a header row.")
    schema = infer_schema(headers, rows)
    dataset_id = str(uuid.uuid4())
    table_name = f"dataset_{uuid.uuid4().hex}"
    table = quote_identifier(table_name)
    definitions = ['"_record_id" VARCHAR PRIMARY KEY']
    definitions.extend(f"{quote_identifier(column['name'])} {DUCK_TYPES[column['type']]}" for column in schema)
    columns = [column["name"] for column in schema]
    connection = None
    with DB_LOCK:
        try:
            connection = connect()
            connection.execute("BEGIN TRANSACTION")
            connection.execute(f"CREATE TABLE {table} ({', '.join(definitions)})")
            if rows:
                placeholders = ", ".join("?" for _ in range(len(columns) + 1))
                insert = f"INSERT INTO {table} VALUES ({placeholders})"
                values = []
                for row in rows:
                    values.append([str(uuid.uuid4()), *[_coerce_cell(row[index] if index < len(row) else None, schema[index]["type"]) for index in range(len(columns))]])
                connection.executemany(insert, values)
            connection.execute(
                "INSERT INTO _datasets (id, name, description, source_type, source_filename, table_name, schema_json, row_count) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                [dataset_id, name.strip(), description.strip(), source_type, filename, table_name, json.dumps(schema), len(rows)],
            )
            connection.execute("COMMIT")
            return _dataset(connection, dataset_id)
        except Exception as error:
            if connection:
                try:
                    connection.execute("ROLLBACK")
                except Exception:
                    pass
            if isinstance(error, HTTPException):
                raise
            raise HTTPException(status_code=422, detail=f"Could not import this dataset: {error}") from error
        finally:
            if connection:
                connection.close()


def _coerce_cell(value: Any, data_type: str):
    if value is None or value == "":
        return None
    if data_type == "text":
        return str(value)
    if data_type == "integer":
        return int(value)
    if data_type == "boolean":
        if isinstance(value, bool):
            return value
        normalized = str(value).strip().casefold()
        if normalized not in {"true", "false"}:
            raise ValueError("expected true or false")
        return normalized == "true"
    if data_type == "number":
        number = float(value)
        if not math.isfinite(number):
            raise ValueError("number must be finite")
        return number
    if data_type == "date":
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        return date.fromisoformat(str(value)[:10])
    if data_type == "datetime":
        if isinstance(value, datetime):
            return value.replace(tzinfo=None)
        return datetime.fromisoformat(str(value).replace("Z", "+00:00")).replace(tzinfo=None)
    return value


def _column(dataset, name: str):
    for column in dataset["schema"]:
        if column["name"] == name:
            return column
    raise HTTPException(status_code=422, detail="A requested column is not part of this dataset.")


def _filter_sql(dataset, definition: dict | None, params: list, depth: int = 0, budget: dict | None = None) -> str:
    if not definition:
        return "TRUE"
    budget = budget if budget is not None else {"remaining": 200}
    budget["remaining"] -= 1
    if budget["remaining"] < 0:
        raise HTTPException(status_code=422, detail="Filters may contain no more than 200 conditions.")
    if not isinstance(definition, dict):
        raise HTTPException(status_code=422, detail="Each filter must be a condition or group.")
    if depth > 5:
        raise HTTPException(status_code=422, detail="Filters may be nested no more than five levels.")
    node_type = definition.get("type", "group" if "conditions" in definition else "condition")
    if node_type == "group":
        logic = str(definition.get("logic", "and")).lower()
        if logic not in {"and", "or"}:
            raise HTTPException(status_code=422, detail="Filter groups must use AND or OR.")
        children = definition.get("conditions", [])
        if not isinstance(children, list) or len(children) > 100:
            raise HTTPException(status_code=422, detail="A filter group must contain at most 100 conditions.")
        compiled = [_filter_sql(dataset, child, params, depth + 1, budget) for child in children]
        compiled = [child for child in compiled if child != "TRUE"]
        return "(" + (" AND " if logic == "and" else " OR ").join(compiled) + ")" if compiled else "TRUE"
    if node_type != "condition":
        raise HTTPException(status_code=422, detail="Filter node types must be condition or group.")
    column = _column(dataset, str(definition.get("column", "")))
    data_type = column["type"]
    operator = definition.get("operator")
    if not isinstance(operator, str) or operator not in FILTER_OPERATORS[data_type]:
        raise HTTPException(status_code=422, detail=f"{operator!s} is not supported for {data_type} columns.")
    field = quote_identifier(column["name"])
    if operator == "is_null":
        return f"{field} IS NULL"
    if operator == "is_not_null":
        return f"{field} IS NOT NULL"
    if operator == "is_empty":
        return f"({field} IS NULL OR {field} = '')"
    if operator == "is_not_empty":
        return f"{field} IS NOT NULL AND {field} <> ''"
    value = definition.get("value")
    if value is None and operator not in {"equals", "not_equals"}:
        raise HTTPException(status_code=422, detail="This filter requires a value.")
    if operator in {"equals", "not_equals"} and value is None:
        return f"{field} IS {'NOT ' if operator == 'not_equals' else ''}NULL"
    try:
        if operator == "between":
            if not isinstance(value, list) or len(value) != 2:
                raise ValueError("between filters require exactly two values")
            if any(item is None or item == "" for item in value):
                raise ValueError("between filters require two values")
            params.extend([_coerce_cell(value[0], data_type), _coerce_cell(value[1], data_type)])
            return f"{field} BETWEEN ? AND ?"
        if operator == "contains":
            params.append(str(value))
            return f"POSITION(? IN CAST({field} AS VARCHAR)) > 0"
        if operator == "not_contains":
            params.append(str(value))
            return f"POSITION(? IN CAST({field} AS VARCHAR)) = 0"
        if operator == "starts_with":
            params.append(str(value))
            return f"starts_with(CAST({field} AS VARCHAR), ?)"
        if operator == "ends_with":
            params.append(str(value))
            return f"ends_with(CAST({field} AS VARCHAR), ?)"
        params.append("" if value == "" and data_type == "text" else _coerce_cell(value, data_type))
    except (ValueError, TypeError) as error:
        raise HTTPException(status_code=422, detail=f"Invalid value for {column['name']}: {error}") from error
    comparisons = {
        "equals": "=",
        "not_equals": "<>",
        "greater_than": ">",
        "greater_or_equal": ">=",
        "less_than": "<",
        "less_or_equal": "<=",
        "on": "=",
        "before": "<",
        "after": ">",
    }
    return f"{field} {comparisons[operator]} ?"


def _where(dataset, search: str, filters: dict | None, params: list) -> str:
    clauses = []
    if search:
        search_clauses = []
        for column in dataset["schema"]:
            search_clauses.append(f"CAST({quote_identifier(column['name'])} AS VARCHAR) ILIKE ?")
            params.append(f"%{search}%")
        if search_clauses:
            clauses.append("(" + " OR ".join(search_clauses) + ")")
    if filters:
        clause = _filter_sql(dataset, filters, params)
        if clause != "TRUE":
            clauses.append(clause)
    return " WHERE " + " AND ".join(clauses) if clauses else ""


def query_records(dataset_id: str, page: int, page_size: int, search: str, sort_by: str | None, sort_dir: str, filters: dict | None):
    page = max(page, 1)
    page_size = min(max(page_size, 1), 200)
    with DB_LOCK:
        connection = connect()
        try:
            dataset = _dataset(connection, dataset_id)
            params: list[Any] = []
            where = _where(dataset, search, filters, params)
            count = connection.execute(f"SELECT count(*) FROM {quote_identifier(dataset['table_name'])}{where}", params).fetchone()[0]
            ordering = ""
            if sort_by:
                _column(dataset, sort_by)
                direction = "DESC" if sort_dir.lower() == "desc" else "ASC"
                ordering = f" ORDER BY {quote_identifier(sort_by)} {direction} NULLS LAST"
            elif dataset["schema"]:
                ordering = f" ORDER BY {quote_identifier(dataset['schema'][0]['name'])} ASC NULLS LAST"
            offset = (page - 1) * page_size
            rows = connection.execute(
                f"SELECT * FROM {quote_identifier(dataset['table_name'])}{where}{ordering} LIMIT ? OFFSET ?",
                [*params, page_size, offset],
            ).fetchall()
            names = ["_record_id", *[column["name"] for column in dataset["schema"]]]
            return {"items": [dict(zip(names, row)) for row in rows], "total": count, "page": page, "page_size": page_size}
        finally:
            connection.close()


def get_record(dataset_id: str, record_id: str):
    with DB_LOCK:
        connection = connect()
        try:
            dataset = _dataset(connection, dataset_id)
            result = connection.execute(
                f"SELECT * FROM {quote_identifier(dataset['table_name'])} WHERE _record_id = ?",
                [record_id],
            ).fetchone()
            if result is None:
                raise HTTPException(status_code=404, detail="Record not found.")
            return dict(zip(["_record_id", *[column["name"] for column in dataset["schema"]]], result))
        finally:
            connection.close()


def save_record(dataset_id: str, values: dict, record_id: str | None = None):
    with DB_LOCK:
        connection = connect()
        try:
            dataset = _dataset(connection, dataset_id)
            columns = dataset["schema"]
            if not isinstance(values, dict) or not values:
                raise HTTPException(status_code=422, detail="Provide at least one field value.")
            unknown = set(values) - {column["name"] for column in columns}
            if unknown:
                raise HTTPException(status_code=422, detail="One or more fields are not part of this dataset.")
            normalized = {}
            try:
                for column in columns:
                    if column["name"] in values:
                        normalized[column["name"]] = _coerce_cell(values[column["name"]], column["type"])
            except (ValueError, TypeError) as error:
                raise HTTPException(status_code=422, detail=f"Invalid field value: {error}") from error
            table = quote_identifier(dataset["table_name"])
            if record_id:
                exists = connection.execute(
                    f"SELECT 1 FROM {table} WHERE _record_id = ?", [record_id]
                ).fetchone()
                if exists is None:
                    raise HTTPException(status_code=404, detail="Record not found.")
                assignments = ", ".join(f"{quote_identifier(name)} = ?" for name in normalized)
                cursor = connection.execute(
                    f"UPDATE {table} SET {assignments} WHERE _record_id = ?",
                    [*normalized.values(), record_id],
                )
                connection.execute("UPDATE _datasets SET updated_at = current_timestamp WHERE id = ?", [dataset_id])
                return get_record(dataset_id, record_id)
            new_id = str(uuid.uuid4())
            insert_values = {"_record_id": new_id, **normalized}
            fields = ", ".join(quote_identifier(name) for name in insert_values)
            placeholders = ", ".join("?" for _ in insert_values)
            connection.execute(f"INSERT INTO {table} ({fields}) VALUES ({placeholders})", list(insert_values.values()))
            connection.execute("UPDATE _datasets SET row_count = row_count + 1, updated_at = current_timestamp WHERE id = ?", [dataset_id])
            return get_record(dataset_id, new_id)
        finally:
            connection.close()


def delete_record(dataset_id: str, record_id: str):
    with DB_LOCK:
        connection = connect()
        try:
            dataset = _dataset(connection, dataset_id)
            count = connection.execute(
                f"SELECT count(*) FROM {quote_identifier(dataset['table_name'])} WHERE _record_id = ?",
                [record_id],
            ).fetchone()[0]
            if not count:
                raise HTTPException(status_code=404, detail="Record not found.")
            connection.execute(f"DELETE FROM {quote_identifier(dataset['table_name'])} WHERE _record_id = ?", [record_id])
            connection.execute("UPDATE _datasets SET row_count = row_count - 1, updated_at = current_timestamp WHERE id = ?", [dataset_id])
            return {"deleted": True}
        finally:
            connection.close()


def save_configuration(dataset_id: str, kind: str, name: str, description: str, config: dict, configuration_id: str | None = None):
    if kind not in {"filter", "table", "detail"}:
        raise HTTPException(status_code=422, detail="Configuration type must be filter, table, or detail.")
    if not name.strip():
        raise HTTPException(status_code=422, detail="Configuration names cannot be empty.")
    with DB_LOCK:
        connection = connect()
        try:
            _dataset(connection, dataset_id)
            config_id = configuration_id or str(uuid.uuid4())
            if configuration_id:
                existing = connection.execute(
                    "SELECT kind FROM _configurations WHERE id = ? AND dataset_id = ?",
                    [configuration_id, dataset_id],
                ).fetchone()
                if existing is None:
                    raise HTTPException(status_code=404, detail="Configuration not found.")
                if existing[0] != kind:
                    raise HTTPException(status_code=422, detail="A saved configuration cannot change its type.")
                connection.execute(
                    "UPDATE _configurations SET name = ?, description = ?, config_json = ?, updated_at = current_timestamp "
                    "WHERE id = ? AND dataset_id = ? AND kind = ?",
                    [name.strip(), description.strip(), json.dumps(config), configuration_id, dataset_id, kind],
                )
            else:
                connection.execute(
                    "INSERT INTO _configurations (id, dataset_id, kind, name, description, config_json) VALUES (?, ?, ?, ?, ?, ?)",
                    [config_id, dataset_id, kind, name.strip(), description.strip(), json.dumps(config)],
                )
            return get_configuration(dataset_id, config_id)
        finally:
            connection.close()


def list_configurations(dataset_id: str, kind: str | None = None):
    with DB_LOCK:
        connection = connect()
        try:
            _dataset(connection, dataset_id)
            if kind:
                rows = connection.execute(
                    "SELECT id, dataset_id, kind, name, description, config_json, schema_version, created_at, updated_at "
                    "FROM _configurations WHERE dataset_id = ? AND kind = ? ORDER BY name",
                    [dataset_id, kind],
                ).fetchall()
            else:
                rows = connection.execute(
                    "SELECT id, dataset_id, kind, name, description, config_json, schema_version, created_at, updated_at "
                    "FROM _configurations WHERE dataset_id = ? ORDER BY name",
                    [dataset_id],
                ).fetchall()
            return [_configuration_dict(row) for row in rows]
        finally:
            connection.close()


def _configuration_dict(row):
    return dict(zip(
        ("id", "dataset_id", "kind", "name", "description", "config_json", "schema_version", "created_at", "updated_at"),
        [*row[:5], decode_json(row[5]), *row[6:]],
    ))


def get_configuration(dataset_id: str, configuration_id: str):
    with DB_LOCK:
        connection = connect()
        try:
            row = connection.execute(
                "SELECT id, dataset_id, kind, name, description, config_json, schema_version, created_at, updated_at "
                "FROM _configurations WHERE id = ? AND dataset_id = ?",
                [configuration_id, dataset_id],
            ).fetchone()
            if row is None:
                raise HTTPException(status_code=404, detail="Configuration not found.")
            return _configuration_dict(row)
        finally:
            connection.close()


def delete_configuration(dataset_id: str, configuration_id: str):
    with DB_LOCK:
        connection = connect()
        try:
            result = connection.execute(
                "DELETE FROM _configurations WHERE id = ? AND dataset_id = ? RETURNING id",
                [configuration_id, dataset_id],
            ).fetchone()
            if result is None:
                raise HTTPException(status_code=404, detail="Configuration not found.")
            return {"deleted": True}
        finally:
            connection.close()


def export_records(dataset_id: str, search: str, filters: dict | None, selected_columns: list[str] | None, scope: str, format: str):
    with DB_LOCK:
        connection = connect()
        try:
            dataset = _dataset(connection, dataset_id)
            columns = selected_columns or [item["name"] for item in dataset["schema"]]
            for name in columns:
                _column(dataset, name)
            params: list[Any] = []
            where = _where(dataset, search if scope == "filtered" else "", filters if scope == "filtered" else None, params)
            cursor = connection.execute(
                f"SELECT {', '.join(quote_identifier(name) for name in columns)} "
                f"FROM {quote_identifier(dataset['table_name'])}{where}",
                params,
            )
            rows = cursor.fetchall()
            if format == "csv":
                output = io.StringIO(newline="")
                writer = csv.writer(output)
                writer.writerow(columns)
                writer.writerows(rows)
                return output.getvalue().encode("utf-8-sig"), "text/csv", ".csv"
            workbook = Workbook(write_only=True)
            sheet = workbook.create_sheet("Data")
            sheet.append(columns)
            for row in rows:
                sheet.append([value.isoformat() if isinstance(value, (date, datetime)) else value for value in row])
            output = io.BytesIO()
            workbook.save(output)
            return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xlsx"
        finally:
            connection.close()
