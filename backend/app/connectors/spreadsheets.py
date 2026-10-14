import csv
import io
import math
import re
from datetime import date, datetime
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}


def _unique_headers(values: list[Any]) -> list[str]:
    headers = []
    seen: set[str] = set()
    for index, value in enumerate(values):
        base = str(value).strip() if value is not None else ""
        base = re.sub(r"\s+", " ", base) or f"Column {index + 1}"
        candidate = base
        suffix = 2
        while candidate.casefold() in seen:
            candidate = f"{base}_{suffix}"
            suffix += 1
        seen.add(candidate.casefold())
        headers.append(candidate)
    return headers


def _parse_csv(data: bytes) -> tuple[list[str], list[list[Any]], list[str]]:
    text = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise HTTPException(status_code=400, detail="Could not decode this CSV file.")
    try:
        sample = text[:8192]
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        rows = list(csv.reader(io.StringIO(text, newline=""), dialect, strict=True))
    except csv.Error as error:
        raise HTTPException(status_code=400, detail="This CSV file contains malformed rows.") from error
    if not rows:
        return [], [], []
    headers = _unique_headers(rows[0])
    width = len(headers)
    normalized = []
    for row in rows[1:]:
        if len(row) > width and any(value != "" for value in row[width:]):
            raise HTTPException(status_code=400, detail="A CSV row contains more cells than the header.")
        values = (row + [None] * width)[:width]
        normalized.append([value if value != "" else None for value in values])
    return headers, normalized, []


def _parse_xlsx(data: bytes, worksheet: str | None) -> tuple[list[str], list[list[Any]], list[str]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as error:
        raise HTTPException(status_code=400, detail="This Excel workbook is malformed or unsupported.") from error
    try:
        sheet_names = workbook.sheetnames
        if worksheet and worksheet not in sheet_names:
            raise HTTPException(status_code=400, detail="The selected worksheet does not exist.")
        sheet = workbook[worksheet] if worksheet else workbook[sheet_names[0]]
        values = list(sheet.iter_rows(values_only=True))
        while values and not any(value is not None for value in values[-1]):
            values.pop()
        if not values:
            return [], [], sheet_names
        headers = _unique_headers(list(values[0]))
        width = len(headers)
        rows = []
        for row in values[1:]:
            if len(row) > width and any(cell is not None and cell != "" for cell in row[width:]):
                raise HTTPException(status_code=400, detail="A worksheet row contains values beyond the header.")
            cells = list(row[:width]) + [None] * max(0, width - len(row))
            rows.append([cell if cell != "" else None for cell in cells])
        return headers, rows, sheet_names
    finally:
        workbook.close()


def read_spreadsheet(filename: str, data: bytes, worksheet: str | None = None):
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Files must be 25 MB or smaller.")
    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in SUPPORTED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Upload a CSV or XLSX file.")
    if extension == ".csv":
        headers, rows, sheets = _parse_csv(data)
    else:
        headers, rows, sheets = _parse_xlsx(data, worksheet)
    if len(headers) > 500:
        raise HTTPException(status_code=400, detail="Files may contain no more than 500 columns.")
    if len(rows) > 1_000_000:
        raise HTTPException(status_code=400, detail="Files may contain no more than 1,000,000 records.")
    return headers, rows, sheets


def _cell_type(value: Any) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "number"
    if isinstance(value, str):
        normalized = value.strip()
        if not normalized:
            return "null"
        if normalized.casefold() in {"true", "false"}:
            return "boolean"
        try:
            date.fromisoformat(normalized)
            return "date"
        except ValueError:
            pass
        try:
            int(normalized)
            return "integer"
        except ValueError:
            try:
                float(normalized)
                return "number"
            except ValueError:
                pass
        try:
            datetime.fromisoformat(normalized.replace("Z", "+00:00"))
            return "datetime"
        except ValueError:
            return "text"
    return "text"


def infer_schema(headers: list[str], rows: list[list[Any]]) -> list[dict[str, str]]:
    schema = []
    for index, name in enumerate(headers):
        found = {_cell_type(row[index]) for row in rows if index < len(row)}
        found.discard("null")
        if not found:
            data_type = "text"
        elif found <= {"integer"}:
            data_type = "integer"
        elif found <= {"integer", "number"}:
            data_type = "number"
        elif found <= {"date", "datetime"}:
            data_type = "datetime" if "datetime" in found else "date"
        elif len(found) == 1:
            data_type = next(iter(found))
        else:
            data_type = "text"
        schema.append({"name": name, "type": data_type})
    return schema
